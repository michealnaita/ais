from enum import Enum

import mammoth
from bs4 import BeautifulSoup
from docx import Document

MAX_CHARS = 2048


class BlockType(str, Enum):
    LIST = "List Paragraph"
    TABLE = "Table"
    HEADING_1 = "Heading 1"
    HEADING_2 = "Heading 2"
    HEADING_3 = "Heading 3"
    CAPTION = "Caption"
    TITLE = "Title"
    NORMAL = "Normal"


style_map = """
p[style-name='Caption'] => figcaption:fresh
"""


class DocumentParser:
    def to_html(self, file_path):
        with open(file_path, "rb") as file:
            result = mammoth.convert_to_html(file, style_map=style_map)
            html = result.value
        soup = BeautifulSoup(html)

    @staticmethod
    def get_chunks(file_path: str):
        chunks = []
        content = Document(file_path)
        blocks_len = len(content.paragraphs)
        for i in range(blocks_len):
            block = content.paragraphs[i]
            block_type = block.style.name
            block_content = block.text

            if len(chunks) == 0:
                return chunks.append({block_type: block_type, "content": block_content})

            if block_type in [BlockType.LIST, BlockType.NORMAL]:
                last_chunk = chunks[-1]
                last_chunk_size = len(last_chunk["content"])
                new_chunk_size = last_chunk_size + len(block_content)
                if (
                    last_chunk["block_type"] == block_type
                    and new_chunk_size <= MAX_CHARS
                ):
                    updated_content = last_chunk["content"] + "\n" + block_content
                    chunks[-1] = {
                        "block_type": block_type,
                        "content": updated_content,
                    }  # replace with last_chunk with a new chunk containing the concatenated content
                else:
                    chunks.append({"block_type": block_type, "content": block_content})
            elif block_type in [
                BlockType.HEADING_1,
                BlockType.HEADING_2,
                BlockType.HEADING_3,
                BlockType.CAPTION,
                BlockType.TITLE,
                BlockType.NORMAL,
            ]:
                chunks.append({block_type: block_type, "content": block_content})
